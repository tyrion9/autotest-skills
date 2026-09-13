#!/usr/bin/env python3
"""Sinh testcase-pairwise.xlsx từ 1 PICT model (+ tuỳ chọn factor.md + gherkin-template).

Triết lý: phần TÍNH TOÁN tổ hợp (pict-cli) và GHÉP CHUỖI (string template) đều
là code xác định — không để AI tự bịa dòng nào. Phần cần phán đoán ngôn ngữ tự
nhiên (cách diễn đạt Gherkin cho tính năng cụ thể) do người/AI làm 1 LẦN khi
viết file --gherkin-template, không làm lại cho từng dòng.

Cách dùng:
    python pict_to_xlsx.py MODEL.txt \
        --gherkin-template TEMPLATE.txt \
        [--factor FACTOR.md] \
        [--order 2] [--seed 42] \
        [-o testcase-pairwise.xlsx] [--check]

MODEL.txt        : file model PICT (xem skill design-pairwise-tests).
TEMPLATE.txt      : mỗi dòng là 1 step Gherkin ĐÃ CÓ từ khoá (When/And...),
                    placeholder dùng ĐÚNG cú pháp Gherkin Scenario Outline
                    <TenFactor> (khớp chính xác tên cột PICT model, phân biệt
                    hoa/thường) — vd: `When khách chọn đồ uống "<Drink>"`.
                    File này được TÁI DÙNG NGUYÊN VẸN ở bước autotest-gen-test
                    (xlsx_to_feature.py) khi dựng khối Scenario Outline.
FACTOR.md         : nếu có, đọc mục "## 4. Case biên / Negative" (bảng markdown)
                    để thêm các dòng Loai=Boundary vào cuối (không qua pict-cli).
--check           : KHÔNG ghi file. Sinh lại trong bộ nhớ rồi so sánh với file
                    -o đang có; khác nhau -> in diff và exit 1. Dùng cho CI để
                    phát hiện xlsx bị sửa tay hoặc model/template đã đổi mà
                    chưa regenerate.

Cột output: STT | MatrixID | <factor...> | Loai | Gherkin | KetQuaMongDoi | GhiChu
  - STT      : số thứ tự để đọc cho tiện (KHÔNG dùng để truy vết).
  - MatrixID : ID ỔN ĐỊNH, hash theo NỘI DUNG tổ hợp (TC-xxxxxx / BC-xxxxxx).
               Cùng 1 tổ hợp giá trị luôn cho cùng 1 ID dù đổi vị trí dòng hay
               sinh lại file -> bug report cũ trích MatrixID vẫn trỏ đúng tổ
               hợp (khác hẳn ID theo số thứ tự, bị xô lệch mỗi lần regenerate).
Sheet "meta" ghi lại provenance (model/seed/order/hash/thống kê pict) để audit.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font
except ImportError:  # pragma: no cover - môi trường thiếu dependency
    print("Cần cài openpyxl trước: pip install openpyxl", file=sys.stderr)
    raise

PICT_SPEC = "pict-cli@0.2"
DATA_SHEET = "testcase-pairwise"
META_SHEET = "meta"

# Tên cột trong bảng "Case biên" của factor.md — chấp nhận nhiều biến thể để
# không phụ thuộc cứng 1 cách viết. Nếu KHÔNG khớp cái nào -> báo lỗi rõ ràng
# (KHÔNG im lặng trả rỗng, vì như vậy case biên sẽ mất dữ liệu mà không ai biết).
DESC_ALIASES = ("Mô tả case", "Mô tả", "Case", "Mo ta case", "Description", "Scenario")
INPUT_ALIASES = ("Input", "Đầu vào", "Dau vao", "Dữ liệu vào")
EXPECTED_ALIASES = ("Kết quả mong đợi", "Kỳ vọng", "Ket qua mong doi", "Expected", "Expected result")

_PLACEHOLDER_RE = re.compile(r"<([A-Za-z0-9_]+)>")


class PictToXlsxError(RuntimeError):
    """Lỗi nghiệp vụ của script — in message gọn, không dump traceback."""


# --------------------------------------------------------------------------- pict-cli


def run_pict(model_path: Path, order: int, seed: int) -> list[dict]:
    if not shutil.which("npx"):
        raise PictToXlsxError(
            "Không tìm thấy 'npx' trong PATH — cần Node.js 22 hoặc 24 để chạy pict-cli."
        )
    cmd = ["npx", "-y", PICT_SPEC, str(model_path), "-o", str(order), "-r", str(seed)]
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        raise PictToXlsxError(
            f"pict-cli lỗi (exit {result.returncode}) với model {model_path}:\n{result.stderr.strip()}"
        )
    lines = [line for line in result.stdout.splitlines() if line.strip()]
    if len(lines) < 2:
        raise PictToXlsxError(
            f"pict-cli không sinh được dòng nào từ {model_path}. stdout:\n{result.stdout}"
        )
    header = lines[0].split("\t")
    return [dict(zip(header, line.split("\t"))) for line in lines[1:]]


def run_pict_stats(model_path: Path, order: int, seed: int) -> str:
    """Chạy pict-cli -s lấy thống kê coverage (Combinations/Generated tests)."""
    cmd = ["npx", "-y", PICT_SPEC, str(model_path), "-o", str(order), "-r", str(seed), "-s"]
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        return "(không lấy được thống kê)"
    return " | ".join(line.strip() for line in result.stdout.splitlines() if line.strip())


# --------------------------------------------------------------------------- ID ổn định


def _sha1(text: str) -> str:
    return hashlib.sha1(text.encode("utf-8")).hexdigest()


def stable_id(prefix: str, payload: str, taken: set[str], length: int = 6) -> str:
    """ID ổn định theo NỘI DUNG (không theo vị trí dòng).

    Kéo dài hash khi đụng độ để đảm bảo duy nhất trong cùng 1 file.
    """
    digest = _sha1(payload).upper()
    size = length
    while size <= len(digest):
        candidate = f"{prefix}-{digest[:size]}"
        if candidate not in taken:
            taken.add(candidate)
            return candidate
        size += 2
    raise PictToXlsxError(f"Không sinh được MatrixID duy nhất cho: {payload!r}")


def pairwise_payload(row: dict) -> str:
    return "|".join(f"{k}={row[k]}" for k in sorted(row))


# --------------------------------------------------------------------------- Gherkin


def render_gherkin(template_lines: list[str], row: dict, allow_unknown: bool = False) -> str:
    """Thay <TenFactor> bằng giá trị thật của dòng.

    Placeholder KHÔNG khớp tên cột nào -> lỗi rõ ràng (thường là gõ sai tên
    factor). Nếu template cố tình chứa `<...>` không phải placeholder (vd thẻ
    HTML trong text), dùng --allow-unknown-placeholders để giữ nguyên.
    """
    rendered = []
    unknown: list[str] = []

    def _sub(match: re.Match) -> str:
        name = match.group(1)
        if name in row:
            return row[name]
        unknown.append(name)
        return match.group(0)

    for line in template_lines:
        rendered.append(_PLACEHOLDER_RE.sub(_sub, line))

    if unknown and not allow_unknown:
        raise PictToXlsxError(
            f"Placeholder không khớp cột nào trong PICT model: {sorted(set(unknown))}.\n"
            f"Các cột hiện có: {sorted(row)}.\n"
            "Sửa lại tên trong --gherkin-template, hoặc thêm --allow-unknown-placeholders "
            "nếu đó là text thật (vd thẻ HTML) chứ không phải placeholder."
        )
    return "\n".join(rendered)


# --------------------------------------------------------------------------- factor.md


def _pick_column(headers: list[str], aliases: tuple[str, ...]) -> str | None:
    normalized = {h.strip().lower(): h for h in headers}
    for alias in aliases:
        if alias.strip().lower() in normalized:
            return normalized[alias.strip().lower()]
    return None


def parse_boundary_rows(factor_md_path: Path) -> list[dict]:
    """Đọc bảng markdown trong mục '## 4. Case biên' của factor.md.

    Trả về list dict đã chuẩn hoá key: desc / input / expected.
    Fail loud nếu có bảng nhưng không nhận ra cột mô tả — tránh tình trạng
    case biên bị bỏ qua âm thầm.
    """
    text = factor_md_path.read_text(encoding="utf-8")
    match = re.search(r"##\s*4\..*?\n(.*?)(\n##\s|\Z)", text, re.S)
    if not match:
        print(
            f"[cảnh báo] {factor_md_path} không có mục '## 4. Case biên' — bỏ qua dòng Boundary.",
            file=sys.stderr,
        )
        return []

    table_lines = [line for line in match.group(1).splitlines() if line.strip().startswith("|")]
    if len(table_lines) < 2:
        print(
            f"[cảnh báo] Mục '## 4. Case biên' trong {factor_md_path} không có bảng dữ liệu — "
            "bỏ qua dòng Boundary.",
            file=sys.stderr,
        )
        return []

    headers = [c.strip() for c in table_lines[0].strip("|").split("|")]
    desc_col = _pick_column(headers, DESC_ALIASES)
    if desc_col is None:
        raise PictToXlsxError(
            f"Bảng 'Case biên' trong {factor_md_path} không có cột mô tả case.\n"
            f"Các cột đang có: {headers}\n"
            f"Cần 1 trong các tên: {list(DESC_ALIASES)}"
        )
    input_col = _pick_column(headers, INPUT_ALIASES)
    expected_col = _pick_column(headers, EXPECTED_ALIASES)
    if expected_col is None:
        print(
            f"[cảnh báo] Bảng 'Case biên' trong {factor_md_path} không có cột kết quả mong đợi "
            f"(một trong {list(EXPECTED_ALIASES)}) — cột KetQuaMongDoi sẽ để trống.",
            file=sys.stderr,
        )

    rows = []
    for line in table_lines[2:]:  # bỏ header + dòng phân cách "---"
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) != len(headers):
            print(
                f"[cảnh báo] Bỏ qua dòng sai số cột trong bảng Case biên: {line.strip()}",
                file=sys.stderr,
            )
            continue
        record = dict(zip(headers, cells))
        desc = record[desc_col]
        if not desc:
            continue
        rows.append(
            {
                "desc": desc,
                "input": record.get(input_col, "") if input_col else "",
                "expected": record.get(expected_col, "") if expected_col else "",
            }
        )
    if not rows:
        print(
            f"[cảnh báo] Bảng 'Case biên' trong {factor_md_path} không có dòng dữ liệu nào.",
            file=sys.stderr,
        )
    return rows


# --------------------------------------------------------------------------- build


def build_rows(args) -> tuple[list[str], list[list], list[list]]:
    """Trả về (header, data_rows, meta_rows) — dựng hoàn toàn trong bộ nhớ để
    dùng chung cho cả chế độ ghi file lẫn --check."""
    pict_rows = run_pict(args.model, args.order, args.seed)
    factor_names = list(pict_rows[0].keys())

    template_lines: list[str] = []
    if args.gherkin_template:
        template_lines = [
            line for line in args.gherkin_template.read_text(encoding="utf-8").splitlines() if line.strip()
        ]

    boundary_rows = parse_boundary_rows(args.factor) if args.factor else []

    header = ["STT", "MatrixID", *factor_names, "Loai", "Gherkin", "KetQuaMongDoi", "GhiChu"]
    taken: set[str] = set()
    data: list[list] = []

    for idx, row in enumerate(pict_rows, start=1):
        matrix_id = stable_id("TC", pairwise_payload(row), taken)
        gherkin = render_gherkin(template_lines, row, args.allow_unknown_placeholders) if template_lines else ""
        data.append([idx, matrix_id, *[row[f] for f in factor_names], "Pairwise", gherkin, "", ""])

    for offset, brow in enumerate(boundary_rows, start=1):
        matrix_id = stable_id("BC", f"{brow['desc']}|{brow['input']}", taken)
        note = f"{brow['desc']}\nInput: {brow['input']}" if brow["input"] else brow["desc"]
        data.append(
            [
                len(pict_rows) + offset,
                matrix_id,
                *["" for _ in factor_names],
                "Boundary",
                note,
                brow["expected"],
                "Viết Gherkin/step tay ở bước gen-test",
            ]
        )

    meta = [
        ["generated_at", datetime.now().isoformat(timespec="seconds")],
        ["model_file", str(args.model)],
        ["model_sha1", _sha1(args.model.read_text(encoding="utf-8"))[:12]],
        ["pict_cli", PICT_SPEC],
        ["order", str(args.order)],
        ["seed", str(args.seed)],
        ["pict_stats", run_pict_stats(args.model, args.order, args.seed)],
        ["factors", ", ".join(factor_names)],
        ["rows_pairwise", str(len(pict_rows))],
        ["rows_boundary", str(len(boundary_rows))],
        [
            "gherkin_template",
            str(args.gherkin_template) if args.gherkin_template else "(không dùng)",
        ],
        [
            "gherkin_template_sha1",
            _sha1(args.gherkin_template.read_text(encoding="utf-8"))[:12] if args.gherkin_template else "",
        ],
        ["factor_file", str(args.factor) if args.factor else "(không dùng)"],
        ["factor_sha1", _sha1(args.factor.read_text(encoding="utf-8"))[:12] if args.factor else ""],
    ]
    return header, data, meta


def write_workbook(path: Path, header: list[str], data: list[list], meta: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in data:
        ws.append(row)

    gherkin_idx = header.index("Gherkin")
    for row_cells in ws.iter_rows(min_row=2):
        row_cells[gherkin_idx].alignment = Alignment(wrap_text=True, vertical="top")
    for col_cells in ws.columns:
        name = col_cells[0].value
        if name == "Gherkin":
            ws.column_dimensions[col_cells[0].column_letter].width = 60
            continue
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)
    ws.freeze_panes = "A2"

    ms = wb.create_sheet(META_SHEET)
    ms.append(["key", "value"])
    for cell in ms[1]:
        cell.font = Font(bold=True)
    for row in meta:
        ms.append(row)
    ms.column_dimensions["A"].width = 24
    ms.column_dimensions["B"].width = 80

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_workbook_data(path: Path) -> tuple[list, list]:
    """Đọc lại phần dữ liệu (bỏ sheet meta) để so sánh ở chế độ --check."""
    wb = load_workbook(path, data_only=True)
    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def check_against_existing(path: Path, header: list[str], data: list[list]) -> int:
    if not path.exists():
        print(f"[CHECK] THIẾU FILE: {path} chưa tồn tại — cần chạy lại script để sinh.", file=sys.stderr)
        return 1

    old_header, old_data = read_workbook_data(path)
    problems: list[str] = []
    if [str(h) for h in old_header] != [str(h) for h in header]:
        problems.append(f"Header khác nhau:\n  file : {old_header}\n  sinh : {header}")

    def norm(rows: list[list]) -> list[tuple]:
        return [tuple("" if c is None else str(c) for c in r) for r in rows]

    old_norm, new_norm = norm(old_data), norm(data)
    if len(old_norm) != len(new_norm):
        problems.append(f"Số dòng khác nhau: file có {len(old_norm)}, sinh ra {len(new_norm)}")

    shown = 0
    for i, (a, b) in enumerate(zip(old_norm, new_norm), start=1):
        if a != b:
            problems.append(f"Dòng {i} khác nhau:\n  file : {a}\n  sinh : {b}")
            shown += 1
            if shown >= 5:
                problems.append("... (còn nữa, chỉ in 5 dòng đầu)")
                break

    if problems:
        print(f"[CHECK] {path} KHÔNG khớp với model/template hiện tại:", file=sys.stderr)
        for p in problems:
            print("  - " + p, file=sys.stderr)
        print(
            "\nNguyên nhân thường gặp: xlsx bị sửa tay, hoặc model/factor/template đã đổi mà "
            "chưa chạy lại script. Chạy lại không kèm --check để regenerate.",
            file=sys.stderr,
        )
        return 1

    print(f"[CHECK] OK — {path} khớp với model/template hiện tại ({len(new_norm)} dòng).")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("model", type=Path, help="File model PICT")
    ap.add_argument("--gherkin-template", type=Path, default=None)
    ap.add_argument("--factor", type=Path, default=None, help="factor.md để lấy case biên")
    ap.add_argument("--order", type=int, default=2, help="Bậc phủ (2=pairwise). Mặc định 2")
    ap.add_argument("--seed", type=int, default=42, help="Seed cho pict-cli -r (tái lập được)")
    ap.add_argument("-o", "--output", type=Path, default=Path("testcase-pairwise.xlsx"))
    ap.add_argument(
        "--allow-unknown-placeholders",
        action="store_true",
        help="Giữ nguyên <...> không khớp tên factor thay vì báo lỗi",
    )
    ap.add_argument(
        "--check",
        action="store_true",
        help="Chỉ so sánh với file -o đang có (dùng cho CI), không ghi đè. Khác -> exit 1",
    )
    args = ap.parse_args()

    for path_arg, label in (
        (args.model, "model"),
        (args.gherkin_template, "--gherkin-template"),
        (args.factor, "--factor"),
    ):
        if path_arg is not None and not path_arg.exists():
            print(f"Không tìm thấy file {label}: {path_arg}", file=sys.stderr)
            return 2

    try:
        header, data, meta = build_rows(args)
    except PictToXlsxError as exc:
        print(f"Lỗi: {exc}", file=sys.stderr)
        return 2

    if args.check:
        return check_against_existing(args.output, header, data)

    write_workbook(args.output, header, data, meta)
    n_pairwise = sum(1 for r in data if r[header.index("Loai")] == "Pairwise")
    n_boundary = len(data) - n_pairwise
    factors = header[2 : header.index("Loai")]
    print(
        f"Đã ghi {args.output} — {n_pairwise} dòng Pairwise + {n_boundary} dòng Boundary "
        f"= {len(data)} dòng, {len(factors)} factor ({', '.join(factors)})"
    )
    print(
        "MatrixID (TC-/BC-) hash theo NỘI DUNG tổ hợp: cùng 1 tổ hợp giá trị luôn ra cùng 1 ID, "
        "kể cả khi đổi vị trí dòng hay file được sinh lại -> bug report cũ vẫn trích đúng."
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
