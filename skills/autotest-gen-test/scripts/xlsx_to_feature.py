#!/usr/bin/env python3
"""Sinh khối Gherkin (.feature) từ testcase-pairwise.xlsx.

testcase-pairwise.xlsx là NGUỒN CHÂN LÝ cho Examples (dòng Loai=Pairwise) và
danh sách case biên (dòng Loai=Boundary). File --gherkin-template (cùng file
đã dùng ở bước autotest-testcase-pairwise, cú pháp <TenFactor>) được tái dùng
NGUYÊN VẸN làm thân bài Scenario Outline — không viết lại, đảm bảo Gherkin
sinh ra khớp 100% với dữ liệu đã render trong cột "Gherkin" của xlsx.

Cách dùng:
    python xlsx_to_feature.py testcase-pairwise.xlsx \
        --gherkin-template <feature>.gherkin-template.txt \
        --feature-name "Tên feature" \
        --scenario-title "Tên Scenario Outline" \
        [--background background.txt] [--then-steps then.txt] \
        [--tag pairwise] [-o out.feature] [--check existing.feature]

Không truyền -o thì in ra stdout để tự chèn/so sánh với .feature hiện có
(khuyến nghị khi cập nhật 1 feature đã tồn tại, tránh ghi đè các Scenario khác
không do script này quản lý).

--check FILE : chế độ CI — KHÔNG ghi gì, chỉ kiểm tra .feature hiện có còn
    đồng bộ với xlsx không (so sánh các step của Scenario Outline + toàn bộ
    dòng bảng Examples). Lệch -> in chi tiết và exit 1.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

DATA_SHEET = "testcase-pairwise"


class XlsxToFeatureError(RuntimeError):
    """Lỗi nghiệp vụ — in message gọn, không dump traceback."""


def read_rows(xlsx_path: Path) -> tuple[list[str], list[dict]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(h) if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        raise XlsxToFeatureError(f"{xlsx_path} rỗng — không có dòng header.") from None
    rows = [dict(zip(header, r)) for r in rows_iter if any(v is not None for v in r)]
    return header, rows


def factor_columns(header: list[str]) -> list[str]:
    """Các cột factor = phần nằm giữa MatrixID và Loai (bỏ STT nếu có)."""
    for required in ("MatrixID", "Loai"):
        if required not in header:
            raise XlsxToFeatureError(
                f"xlsx thiếu cột bắt buộc '{required}'. Header hiện có: {header}\n"
                "File này có đúng là output của pict_to_xlsx.py không?"
            )
    start = header.index("MatrixID") + 1
    end = header.index("Loai")
    return [c for c in header[start:end] if c and c != "STT"]


def escape_cell(value) -> str:
    """Chuẩn hoá 1 giá trị cho ô bảng Gherkin.

    Gherkin dùng '|' làm dấu phân cách cột và không cho xuống dòng trong ô ->
    phải escape '|' và làm phẳng newline, nếu không bảng Examples sẽ vỡ.
    """
    text = "" if value is None else str(value)
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\r", " ").replace("\n", " ").strip()


def build_examples_table(factors: list[str], pairwise_rows: list[dict]) -> list[str]:
    cols = ["MatrixID", *factors]
    lines = ["    Examples:", "      | " + " | ".join(cols) + " |"]
    for row in pairwise_rows:
        lines.append("      | " + " | ".join(escape_cell(row.get(c)) for c in cols) + " |")
    return lines


def indent(lines: list[str], n: int) -> list[str]:
    pad = " " * n
    return [f"{pad}{line}" if line.strip() else line for line in lines]


def read_lines(path: Path | None) -> list[str]:
    if path is None:
        return []
    return [line for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def build_feature(args, factors: list[str], pairwise_rows: list[dict], boundary_rows: list[dict]) -> tuple[list[str], list[str], list[str]]:
    """Trả về (toàn bộ dòng feature, dòng step của Scenario Outline, dòng bảng Examples)."""
    template_lines = read_lines(args.gherkin_template)
    if not template_lines:
        raise XlsxToFeatureError(f"--gherkin-template {args.gherkin_template} rỗng.")
    then_lines = read_lines(args.then_steps)
    background_lines = read_lines(args.background)

    outline_steps = template_lines + then_lines
    examples_lines = build_examples_table(factors, pairwise_rows)

    out: list[str] = [f"Feature: {args.feature_name}", ""]
    if background_lines:
        out += ["  Background:"]
        out += indent(background_lines, 4)
        out.append("")

    out += [f"  @{args.tag}"]
    out += [f"  Scenario Outline: [<MatrixID>] {args.scenario_title}"]
    out += indent(outline_steps, 4)
    out.append("")
    out += examples_lines
    out.append("")

    if boundary_rows:
        out += [
            "  Rule: Kiểm tra dữ liệu biên / không hợp lệ (từ factor.md mục 4, KHÔNG thuộc bảng pairwise)",
            "",
        ]
        for row in boundary_rows:
            note = str(row.get("Gherkin") or "").strip() or "(chưa có mô tả)"
            expected = str(row.get("KetQuaMongDoi") or "").strip()
            matrix_id = row.get("MatrixID")
            first_line = note.splitlines()[0]
            out += [f"    Scenario: [{matrix_id}] {first_line}"]
            out += ["      # TODO: viết cụ thể Given/When/Then cho case này. Tham khảo:"]
            for note_line in note.splitlines():
                out.append(f"      #   {note_line}")
            if expected:
                out.append(f"      #   Kỳ vọng: {expected}")
            else:
                out.append("      #   [cảnh báo] Chưa có 'Kết quả mong đợi' trong xlsx — bổ sung trước khi viết step.")
            out.append("")

    return out, indent(outline_steps, 4), examples_lines


def check_feature(existing: Path, outline_steps: list[str], examples_lines: list[str]) -> int:
    """So sánh .feature hiện có với dữ liệu sinh từ xlsx (step outline + Examples)."""
    if not existing.exists():
        print(f"[CHECK] THIẾU FILE: {existing} chưa tồn tại.", file=sys.stderr)
        return 1

    text = existing.read_text(encoding="utf-8")
    problems: list[str] = []

    for step in outline_steps:
        if step.strip() and step.strip() not in text:
            problems.append(f"Thiếu step của Scenario Outline: {step.strip()}")

    existing_table = [
        line.rstrip() for line in text.splitlines() if line.strip().startswith("|")
    ]
    expected_table = [line.rstrip() for line in examples_lines if line.strip().startswith("|")]

    if existing_table != expected_table:
        problems.append(
            f"Bảng Examples lệch: file có {len(existing_table)} dòng, xlsx sinh ra {len(expected_table)} dòng."
        )
        shown = 0
        for i, (a, b) in enumerate(zip(existing_table, expected_table), start=1):
            if a != b:
                problems.append(f"  dòng bảng {i}:\n    file : {a.strip()}\n    xlsx : {b.strip()}")
                shown += 1
                if shown >= 5:
                    problems.append("  ... (chỉ in 5 dòng đầu)")
                    break

    if problems:
        print(f"[CHECK] {existing} KHÔNG đồng bộ với xlsx:", file=sys.stderr)
        for p in problems:
            print("  - " + p, file=sys.stderr)
        print(
            "\nChạy lại xlsx_to_feature.py (không kèm --check) rồi cập nhật khối Scenario Outline + "
            "Examples trong .feature cho khớp.",
            file=sys.stderr,
        )
        return 1

    print(f"[CHECK] OK — {existing} đồng bộ với xlsx ({len(expected_table) - 1} dòng Examples).")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--gherkin-template", type=Path, required=True)
    ap.add_argument("--feature-name", required=True)
    ap.add_argument("--scenario-title", required=True)
    ap.add_argument("--background", type=Path, default=None, help="File chứa step Given (Background), tuỳ chọn")
    ap.add_argument("--then-steps", type=Path, default=None, help="File chứa step Then/And nối sau template, tuỳ chọn")
    ap.add_argument("--tag", default="pairwise", help="Tag gắn trên Scenario Outline, mặc định 'pairwise'")
    ap.add_argument("-o", "--output", type=Path, default=None, help="Ghi ra file thay vì in stdout")
    ap.add_argument(
        "--check",
        type=Path,
        default=None,
        metavar="FEATURE",
        help="Chế độ CI: kiểm tra .feature này còn đồng bộ với xlsx không (không ghi gì)",
    )
    args = ap.parse_args()

    for path_arg, label in ((args.xlsx, "xlsx"), (args.gherkin_template, "--gherkin-template")):
        if not path_arg.exists():
            print(f"Không tìm thấy file {label}: {path_arg}", file=sys.stderr)
            return 2

    try:
        header, rows = read_rows(args.xlsx)
        factors = factor_columns(header)
        pairwise_rows = [r for r in rows if r.get("Loai") == "Pairwise"]
        boundary_rows = [r for r in rows if r.get("Loai") == "Boundary"]
        if not pairwise_rows:
            raise XlsxToFeatureError(
                f"{args.xlsx} không có dòng nào Loai='Pairwise' — kiểm tra lại file đầu vào."
            )
        feature_lines, outline_steps, examples_lines = build_feature(args, factors, pairwise_rows, boundary_rows)
    except XlsxToFeatureError as exc:
        print(f"Lỗi: {exc}", file=sys.stderr)
        return 2

    if args.check:
        return check_feature(args.check, outline_steps, examples_lines)

    text = "\n".join(feature_lines).rstrip() + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text, encoding="utf-8")
        print(f"Đã ghi {args.output} ({len(pairwise_rows)} dòng pairwise, {len(boundary_rows)} case biên)")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
